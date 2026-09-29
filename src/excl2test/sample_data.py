"""Génère deux fichiers Excel d'exemple : un valide, un volontairement défectueux.

Usage :  py -m excl2test.sample_data
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import schema

# Cas valides : (ID, Exigence, Sig_in, Val_in, Unit_in, Sig_out, Val_att, Unit_out, Tol_%)
VALID_CASES = [
    ("TC-001", "REQ-DCDC-001", "Vin", 48, "V", "Vout", 12.0, "V", 2),
    ("TC-002", "REQ-DCDC-001", "Vin", 36, "V", "Vout", 12.0, "V", 2),
    ("TC-003", "REQ-DCDC-001", "Vin", 60, "V", "Vout", 12.0, "V", 2),
    ("TC-004", "REQ-DCDC-002", "Iload", 1, "A", "Vout", 12.0, "V", 2),
    ("TC-005", "REQ-DCDC-002", "Iload", 5, "A", "Vout", 12.0, "V", 3),
    ("TC-006", "REQ-DCDC-002", "Iload", 10, "A", "Vout", 11.9, "V", 3),
    ("TC-007", "REQ-DCDC-003", "Temp", -40, "degC", "Vout", 12.0, "V", 5),
    ("TC-008", "REQ-DCDC-003", "Temp", 125, "degC", "Vout", 12.0, "V", 5),
    ("TC-009", "REQ-DCDC-004", "Fsw", 100, "kHz", "Vout", 12.0, "V", 2),
    ("TC-010", "REQ-DCDC-004", "Fsw", 500, "kHz", "Vout", 12.0, "V", 2),
]

# Cas défectueux : même forme que ci-dessus. Chaque ligne porte UN défaut précis.
INVALID_CASES = [
    ("TC-101", "REQ-DCDC-001", "Vinn", 48, "V", "Vout", 12.0, "V", 2),         # signal inconnu
    ("TC-102", "", "Vin", 48, "V", "Vout", 12.0, "V", 2),                      # exigence vide
    ("TC-103", "REQ-DCDC-001", "Vin", "abc", "V", "Vout", 12.0, "V", 2),       # valeur non numérique
    ("TC-104", "REQ-DCDC-001", "Vin", 48, "A", "Vout", 12.0, "V", 2),          # unité incohérente (entrée)
    ("TC-105", "REQ-DCDC-001", "Vin", 500, "V", "Vout", 12.0, "V", 2),         # valeur hors plage
    ("TC-106", "REQ-DCDC-002", "Iload", 5, "A", "Vout", 12.0, "V", 150),       # tolérance hors 0-100
    ("TC-107", "REQ-DCDC-002", "Iload", 5, "A", "Vout", None, "V", 2),         # valeur attendue vide
    ("TC1", "REQ-DCDC-003", "Temp", 25, "degC", "Vout", 12.0, "V", 5),         # format d'ID invalide
    ("TC-109", "REQ-DCDC-004", "Fsw", 100, "kHz", "Vout", 12.0, "mA", 2),      # unité incohérente (sortie)
    ("TC-101", "REQ-DCDC-004", "Fsw", 200, "kHz", "Vout", 12.0, "V", 2),       # ID en doublon
]


def _write_header(ws, columns):
    fill = PatternFill("solid", fgColor="1F4E78")
    for col, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        ws.column_dimensions[get_column_letter(col)].width = max(14, len(name) + 4)
    ws.freeze_panes = "A2"


def build_workbook(cases):
    """Construit un classeur avec la feuille de cas de test et le dictionnaire."""
    wb = Workbook()
    ws_cases = wb.active
    ws_cases.title = schema.SHEET_CASES
    _write_header(ws_cases, schema.CASES_COLUMNS)
    for row in cases:
        ws_cases.append(list(row))

    ws_dict = wb.create_sheet(schema.SHEET_DICT)
    _write_header(ws_dict, schema.DICT_COLUMNS)
    for row in schema.DICTIONARY:
        ws_dict.append(list(row))
    return wb


def generate(output_dir):
    """Écrit les deux fichiers d'exemple et renvoie leurs chemins."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    valid_path = output_dir / "cas_valides.xlsx"
    invalid_path = output_dir / "cas_invalides.xlsx"
    build_workbook(VALID_CASES).save(valid_path)
    build_workbook(INVALID_CASES).save(invalid_path)
    return valid_path, invalid_path


if __name__ == "__main__":
    for path in generate("data"):
        print(f"Fichier créé : {path}")
