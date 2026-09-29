"""Lecture du classeur Excel : structure vérifiée, contenu brut, aucune interprétation.

Toute anomalie de structure (fichier illisible, feuille ou colonne manquante)
lève TemplateError : on ne tente jamais de deviner ce que l'utilisateur voulait.
"""

from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from . import schema
from .models import DictEntry, RawCase


class TemplateError(Exception):
    """Le fichier ne respecte pas la structure du template : analyse impossible."""


def _is_blank(value):
    return value is None or (isinstance(value, str) and value.strip() == "")


def _check_header(ws, expected, sheet_name):
    found = [ws.cell(row=1, column=i).value for i in range(1, len(expected) + 1)]
    for index, (want, got) in enumerate(zip(expected, found), start=1):
        if got != want:
            cell = f"{get_column_letter(index)}1"
            raise TemplateError(
                f"Feuille '{sheet_name}', cellule {cell} : en-tête attendu '{want}', trouvé '{got}'."
            )


def _read_dictionary(ws):
    entries = {}
    for row in ws.iter_rows(min_row=2, max_col=len(schema.DICT_COLUMNS)):
        values = [c.value for c in row]
        if all(_is_blank(v) for v in values):
            continue
        signal, unit, vmin, vmax = values[0], values[1], values[2], values[3]
        line = row[0].row
        for label, col, val in (("Min", "C", vmin), ("Max", "D", vmax)):
            if isinstance(val, bool) or not isinstance(val, (int, float)):
                raise TemplateError(
                    f"Feuille '{schema.SHEET_DICT}', cellule {col}{line} : "
                    f"{label} doit être un nombre, trouvé {val!r}."
                )
        if _is_blank(signal) or _is_blank(unit):
            raise TemplateError(
                f"Feuille '{schema.SHEET_DICT}', ligne {line} : signal ou unité vide."
            )
        if signal in entries:
            raise TemplateError(
                f"Feuille '{schema.SHEET_DICT}', ligne {line} : signal '{signal}' défini deux fois."
            )
        entries[signal] = DictEntry(unit=unit, min=vmin, max=vmax)
    return entries


def _read_cases(ws):
    cases = []
    width = len(schema.CASES_COLUMNS)
    for row in ws.iter_rows(min_row=2, max_col=width):
        values = [c.value for c in row]
        if all(_is_blank(v) for v in values):
            continue  # ligne entièrement vide : ignorée
        cases.append(RawCase(row=row[0].row, values=dict(zip(schema.CASES_COLUMNS, values))))
    return cases


def load(path):
    """Lit le fichier et renvoie (cas bruts, dictionnaire de données)."""
    path = Path(path)
    if not path.exists():
        raise TemplateError(f"Fichier introuvable : {path}")
    try:
        wb = load_workbook(path, data_only=True)
    except (InvalidFileException, BadZipFile, OSError) as exc:
        raise TemplateError(f"Fichier illisible (est-ce bien un .xlsx ?) : {exc}") from exc

    for name in (schema.SHEET_CASES, schema.SHEET_DICT):
        if name not in wb.sheetnames:
            raise TemplateError(
                f"Feuille '{name}' absente. Feuilles trouvées : {wb.sheetnames}."
            )

    _check_header(wb[schema.SHEET_CASES], schema.CASES_COLUMNS, schema.SHEET_CASES)
    _check_header(wb[schema.SHEET_DICT], schema.DICT_COLUMNS, schema.SHEET_DICT)
    dictionary = _read_dictionary(wb[schema.SHEET_DICT])
    cases = _read_cases(wb[schema.SHEET_CASES])
    return cases, dictionary
