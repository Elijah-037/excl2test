import pytest
from openpyxl import load_workbook

from excl2test import schema
from excl2test.sample_data import INVALID_CASES, VALID_CASES, generate


@pytest.mark.exigence("EXG-06")
def test_generate_creates_both_files(tmp_path):
    valid_path, invalid_path = generate(tmp_path)
    assert valid_path.exists()
    assert invalid_path.exists()


@pytest.mark.exigence("EXG-01", "EXG-06")
def test_workbook_structure_matches_schema(tmp_path):
    valid_path, _ = generate(tmp_path)
    wb = load_workbook(valid_path)
    assert wb.sheetnames == [schema.SHEET_CASES, schema.SHEET_DICT]
    header = [c.value for c in wb[schema.SHEET_CASES][1]]
    assert header == schema.CASES_COLUMNS


@pytest.mark.exigence("EXG-06")
def test_row_counts(tmp_path):
    valid_path, invalid_path = generate(tmp_path)
    ws_valid = load_workbook(valid_path)[schema.SHEET_CASES]
    ws_invalid = load_workbook(invalid_path)[schema.SHEET_CASES]
    assert ws_valid.max_row - 1 == len(VALID_CASES) == 10
    assert ws_invalid.max_row - 1 == len(INVALID_CASES) == 10
